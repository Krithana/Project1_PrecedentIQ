import io
import datetime
from typing import Optional, List
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from ..schemas.argument import LegalResearchBrief, ProContraOutline
from ..schemas.analysis import HoldingExtraction
from ..schemas.retrieval import ScoredChunk


class WordBriefExporter:
    """Generates court-ready Word (.docx) chamber briefs."""

    @classmethod
    def create_brief_document(cls, brief: LegalResearchBrief) -> io.BytesIO:
        doc = Document()

        # Set page margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

        # Base styles
        normal_style = doc.styles['Normal']
        normal_style.font.name = 'Calibri'
        normal_style.font.size = Pt(11)
        normal_style.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

        # Title Header
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run("PRECEDENTIQ LEGAL RESEARCH BRIEF")
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(0x0F, 0x2B, 0x5C)  # Deep Navy

        subtitle_p = doc.add_paragraph()
        subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = subtitle_p.add_run("Indian Supreme Court & High Court Case Law Analysis")
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

        # Divider line
        cls._add_divider(doc)

        # Metadata Block
        meta_table = doc.add_table(rows=3, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.autofit = False

        rows = meta_table.rows
        rows[0].cells[0].paragraphs[0].add_run("Research Query:").bold = True
        rows[0].cells[1].paragraphs[0].add_run(brief.query)

        rows[1].cells[0].paragraphs[0].add_run("Matter / File:").bold = True
        rows[1].cells[1].paragraphs[0].add_run(brief.case_file_name or "General Precedent Discovery")

        rows[2].cells[0].paragraphs[0].add_run("Generated On:").bold = True
        rows[2].cells[1].paragraphs[0].add_run(brief.timestamp)

        for row in rows:
            row.cells[0].width = Inches(1.8)
            row.cells[1].width = Inches(4.8)

        doc.add_paragraph()

        # Section 1: Executive Summary
        h1 = doc.add_heading("1. Executive Legal Summary", level=1)
        h1.style.font.color.rgb = RGBColor(0x0F, 0x2B, 0x5C)
        outline = brief.argument_outline
        p_summary = doc.add_paragraph(outline.issue_summary)
        p_summary.paragraph_format.line_spacing = 1.15

        if outline.client_context_summary:
            p_ctx = doc.add_paragraph()
            p_ctx.add_run("Client Case Context: ").bold = True
            p_ctx.add_run(outline.client_context_summary)

        # Section 2: Precedent Matrix
        h2 = doc.add_heading("2. Retrieved Authorities & Precedents Matrix", level=1)
        h2.style.font.color.rgb = RGBColor(0x0F, 0x2B, 0x5C)

        if brief.top_precedents:
            table = doc.add_table(rows=1, cols=4)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = 'Table Grid'
            hdr_cells = table.rows[0].cells
            hdr_titles = ["Citation / Case", "Court", "Section", "Relevance"]
            for i, title in enumerate(hdr_titles):
                hdr_cells[i].paragraphs[0].add_run(title).bold = True
                cls._set_cell_background(hdr_cells[i], "0F2B5C")
                hdr_cells[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            seen_judgments = set()
            for sc in brief.top_precedents:
                if sc.chunk.judgment_id in seen_judgments or sc.final_score < 0:
                    continue
                seen_judgments.add(sc.chunk.judgment_id)
                row_cells = table.add_row().cells
                row_cells[0].paragraphs[0].add_run(sc.chunk.citation)
                court_label = "Supreme Court" if sc.chunk.court == "supreme_court" else "High Court"
                row_cells[1].paragraphs[0].add_run(court_label)
                row_cells[2].paragraphs[0].add_run(sc.chunk.section.replace("_", " ").title())
                row_cells[3].paragraphs[0].add_run(f"{sc.final_score:.3f}")

            doc.add_paragraph()

        # Section 3: Detailed Holding Analysis (Ratio Decidendi)
        if brief.holdings:
            h3 = doc.add_heading("3. Holding Extraction (Ratio Decidendi vs Obiter Dicta)", level=1)
            h3.style.font.color.rgb = RGBColor(0x0F, 0x2B, 0x5C)

            for h in brief.holdings:
                p_case = doc.add_paragraph()
                p_case.add_run(f"• Judgment: {h.case_title or h.judgment_id}").bold = True
                p_case.add_run(f" ({h.court.replace('_', ' ').title()})")
                if h.decision_status:
                    p_status = doc.add_paragraph()
                    p_status.paragraph_format.left_indent = Inches(0.25)
                    p_status.add_run("Status: ").bold = True
                    p_status.add_run(h.decision_status)

                # Ratio box
                p_ratio = doc.add_paragraph()
                p_ratio.paragraph_format.left_indent = Inches(0.25)
                run_r = p_ratio.add_run("Candidate Ratio Decidendi (Verify): ")
                run_r.bold = True
                run_r.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)  # Forest Green
                p_ratio.add_run(h.ratio_decidendi)

                if h.ratio_evidence:
                    p_ev = doc.add_paragraph()
                    p_ev.paragraph_format.left_indent = Inches(0.4)
                    p_ev.add_run("Evidence: ").bold = True
                    p_ev.add_run("; ".join(h.ratio_evidence[:2])).italic = True
                if h.ratio_paragraphs:
                    p_ref = doc.add_paragraph()
                    p_ref.paragraph_format.left_indent = Inches(0.4)
                    p_ref.add_run("References: ").bold = True
                    p_ref.add_run(", ".join(h.ratio_paragraphs))

                if h.obiter_dicta:
                    p_obiter = doc.add_paragraph()
                    p_obiter.paragraph_format.left_indent = Inches(0.25)
                    run_o = p_obiter.add_run("Non-ratio Observation (Verify Legal Effect): ")
                    run_o.italic = True
                    p_obiter.add_run(h.obiter_dicta)

                if h.final_order:
                    p_ord = doc.add_paragraph()
                    p_ord.paragraph_format.left_indent = Inches(0.25)
                    p_ord.add_run("Operative Order: ").bold = True
                    p_ord.add_run(h.final_order)
                    if h.order_evidence:
                        p_ord.add_run(" Evidence: ").bold = True
                        p_ord.add_run("; ".join(h.order_evidence[:2])).italic = True

        # Section 4: Pro & Contra Argument Outline
        h4 = doc.add_heading("4. Argument Synthesis (Pro & Contra Outline)", level=1)
        h4.style.font.color.rgb = RGBColor(0x0F, 0x2B, 0x5C)

        doc.add_heading("A. Affirmative Grounds (Supporting Arguments)", level=2)
        if outline.supporting_arguments:
            for pt in outline.supporting_arguments:
                p = doc.add_paragraph()
                p.add_run(f"[{pt.strength.upper()}] {pt.title}: ").bold = True
                p.add_run(pt.legal_proposition)

                p_reason = doc.add_paragraph()
                p_reason.paragraph_format.left_indent = Inches(0.25)
                p_reason.add_run("Application: ").italic = True
                p_reason.add_run(pt.applicability_reasoning)

                for cit in pt.supporting_citations:
                    p_cit = doc.add_paragraph()
                    p_cit.paragraph_format.left_indent = Inches(0.4)
                    p_cit.add_run(f"Citation: {cit.citation} (source chunk {cit.chunk_id})\n").bold = True
                    p_cit.add_run(f"Excerpt: \"{cit.excerpt}\"").italic = True
        else:
            doc.add_paragraph("No affirmative authorities identified.")

        doc.add_heading("B. Anticipated Opposition & Factual Distinctions (Contra Arguments)", level=2)
        if outline.counter_arguments:
            for cpt in outline.counter_arguments:
                p_opp = doc.add_paragraph()
                p_opp.add_run("Potential Objection: ").bold = True
                p_opp.add_run(cpt.potential_objection)

                p_dist = doc.add_paragraph()
                p_dist.paragraph_format.left_indent = Inches(0.25)
                p_dist.add_run("Grounds to Distinguish: ").bold = True
                p_dist.add_run(cpt.rebuttal_or_distinction)

                for cit in cpt.adverse_citations:
                    p_cit = doc.add_paragraph()
                    p_cit.paragraph_format.left_indent = Inches(0.4)
                    p_cit.add_run(f"Adverse Authority: {cit.citation} (source chunk {cit.chunk_id})\n").bold = True
                    p_cit.add_run(f"Excerpt: \"{cit.excerpt}\"").italic = True
        else:
            doc.add_paragraph("No significant adverse authorities identified.")

        # Strategic Advice
        if outline.strategic_advice:
            doc.add_heading("5. Strategic Litigation Advice", level=1)
            p_adv = doc.add_paragraph(outline.strategic_advice)
            p_adv.paragraph_format.left_indent = Inches(0.2)

        # Verification Notice
        cls._add_divider(doc)
        p_foot = doc.add_paragraph()
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        foot_run = p_foot.add_run(
            "Grounded Legal AI Brief. Source excerpts are linked to the repository; legal classifications are candidates. "
            "Independent advocate verification is required before court filing."
        )
        foot_run.font.size = Pt(8.5)
        foot_run.font.italic = True
        foot_run.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

        output = io.BytesIO()
        doc.save(output)
        output.seek(0)
        return output

    @staticmethod
    def _add_divider(doc):
        p = doc.add_paragraph()
        p_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                             r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="0F2B5C"/>'
                             r'</w:pBdr>')
        p._p.get_or_add_pPr().append(p_border)

    @staticmethod
    def _set_cell_background(cell, fill_hex: str):
        tcPr = cell._element.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:val="clear" w:color="auto" w:fill="{fill_hex}"/>')
        tcPr.append(shd)
